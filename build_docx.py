# -*- coding: utf-8 -*-
"""按公文格式要求生成 .docx：
标题二号方正小标宋简体居中；作者三号楷体_GB2312居中；单位邮箱四号宋体居中；
摘要/关键词标签四号黑体、内容四号楷体_GB2312；
一级标题三号黑体、二级三号楷体、三级三号仿宋_GB2312加粗；正文三号仿宋_GB2312；
页码四号宋体居中，格式 —1—。摘要关键词仅中文。
"""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = '/home/user/battle/论文终稿-非战争军事行动中智能无人作战的策略战法研究.md'
OUT = '/home/user/battle/非战争军事行动中智能无人作战的策略战法研究.docx'
FIG = '/home/user/battle/figs/'

# 字体
XBS  = '方正小标宋简体'
KAI  = '楷体_GB2312'
FS   = '仿宋_GB2312'
HEI  = '黑体'
SONG = '宋体'
TIMES = 'Times New Roman'

# 字号（磅）
HAO2, HAO3, HAO4 = 22, 16, 14

# 行距：三号正文用固定值 29 磅（公文常用 28—30）
LINE_BODY = 29


def setfont(run, cjk_font, size, bold=False, western=TIMES):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = western
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), cjk_font)
    rFonts.set(qn('w:ascii'), western)
    rFonts.set(qn('w:hAnsi'), western)


def fixed_line(pf, pts=LINE_BODY):
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(pts)


def straight_to_curly(t):
    out, opening = [], True
    for ch in t:
        if ch == '"':
            out.append('“' if opening else '”')
            opening = not opening
        else:
            out.append(ch)
    return ''.join(out)


def newp(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent_chars=0, size=HAO3,
         before=0, after=0, line=LINE_BODY):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    fixed_line(pf, line)
    if indent_chars:
        pf.first_line_indent = Pt(size * indent_chars)
    return p


def _plain(p, seg, cjk_font, size, bold):
    """处理转义、下标（X_y）与上标星号（p\\*）"""
    seg = seg.replace('\\*', '*')
    for sub in re.split(r'([A-Za-z]_[A-Za-z0-9]+)', seg):
        if not sub:
            continue
        m = re.fullmatch(r'([A-Za-z])_([A-Za-z0-9]+)', sub)
        if m:
            r1 = p.add_run(m.group(1)); setfont(r1, cjk_font, size, bold)
            r1.font.italic = True
            r2 = p.add_run(m.group(2)); setfont(r2, cjk_font, size, bold)
            r2.font.subscript = True
        else:
            r = p.add_run(sub); setfont(r, cjk_font, size, bold)


def rich(p, text, cjk_font, size, bold=False):
    """解析 **粗体**、【n】上标、下标与转义"""
    for seg in re.split(r'(\*\*[^*]+\*\*|【\d+】)', text):
        if not seg:
            continue
        if seg.startswith('**') and seg.endswith('**'):
            _plain(p, seg[2:-2], cjk_font, size, True)
        elif re.fullmatch(r'【\d+】', seg):
            r = p.add_run('[' + seg[1:-1] + ']')
            setfont(r, cjk_font, size)
            r.font.superscript = True
        else:
            _plain(p, seg, cjk_font, size, bold)


def add_picture(doc, fname, caption, width_cm):
    p = newp(doc, WD_ALIGN_PARAGRAPH.CENTER, 0, before=8, after=2, line=0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.add_run().add_picture(FIG + fname, width=Cm(width_cm))
    c = newp(doc, WD_ALIGN_PARAGRAPH.CENTER, 0, before=0, after=10, line=0)
    c.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    r = c.add_run(caption); setfont(r, HEI, 12)


def add_formula(doc, img, number, width_cm):
    p = newp(doc, WD_ALIGN_PARAGRAPH.CENTER, 0, before=6, after=6, line=0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.add_run().add_picture(FIG + img, width=Cm(width_cm))
    r = p.add_run('　　（' + number + '）'); setfont(r, SONG, HAO4)


def add_table(doc, rows, caption):
    c = newp(doc, WD_ALIGN_PARAGRAPH.CENTER, 0, before=8, after=3, line=0)
    c.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    r = c.add_run(caption); setfont(r, HEI, 12, True)
    tbl = doc.add_table(rows=len(rows), cols=len(rows[0]))
    tbl.style = 'Table Grid'
    tbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for ri, row in enumerate(rows):
        for ci, cell in enumerate(row):
            tc = tbl.cell(ri, ci)
            para = tc.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 else WD_ALIGN_PARAGRAPH.LEFT
            pf = para.paragraph_format
            pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
            pf.space_after = Pt(1)
            rr = para.add_run(cell)
            setfont(rr, HEI if ri == 0 else FS, 10.5, bold=(ri == 0))


def add_page_number_footer(section):
    """页码：四号宋体居中，格式 —1—"""
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    for r in list(p.runs):
        r._element.getparent().remove(r._element)

    r1 = p.add_run('—'); setfont(r1, SONG, HAO4)
    # PAGE 域
    rf = p.add_run()
    setfont(rf, SONG, HAO4)
    fb = OxmlElement('w:fldChar'); fb.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = ' PAGE '
    fs_ = OxmlElement('w:fldChar'); fs_.set(qn('w:fldCharType'), 'separate')
    tt = OxmlElement('w:t'); tt.text = '1'
    fe = OxmlElement('w:fldChar'); fe.set(qn('w:fldCharType'), 'end')
    for el in (fb, it, fs_, tt, fe):
        rf._element.append(el)
    r2 = p.add_run('—'); setfont(r2, SONG, HAO4)


# ═════════════ 解析并生成 ═════════════
md = straight_to_curly(open(SRC, encoding='utf-8').read())
lines = md.split('\n')

doc = Document()
sec = doc.sections[0]
# 公文页边距 GB/T 9704：上37 下35 左28 右26 mm
sec.top_margin = Cm(3.7); sec.bottom_margin = Cm(3.5)
sec.left_margin = Cm(2.8); sec.right_margin = Cm(2.6)
add_page_number_footer(sec)

st = doc.styles['Normal']
st.font.name = TIMES; st.font.size = Pt(HAO3)
st.element.rPr.rFonts.set(qn('w:eastAsia'), FS)

FORMULAS = {'FORMULA1': ('formula1.png', '1', 7.0),
            'FORMULA2': ('formula2.png', '2', 3.2),
            'FORMULA3': ('formula3.png', '3', 4.6)}

# 图插入点：(所属标题, 该标题下第 n 个正文段落之后)
FIGS_AFTER = {
    ('一、引言', 2): ('fig1_framework.png', '图1　研究框架与逻辑主线', 14.5),
    ('1.“图—谱—链”三源复合侦测', 2): ('fig2_three_sources.png', '图2　“图—谱—链”三源复合侦测融合', 14.5),
    ('2.“圈层—网格—白名单”分区管控', 1): ('fig3_zoning.png', '图3　“圈层—网格—白名单”分区管控', 12.5),
    ('2. 授权前置与快速响应的张力及其化解', 2): ('fig4_legal_response.png', '图4　法律授权前置的分级响应模型', 14.8),
    ('（二）一体化的代价与适用边界', 1): ('fig5_rescue_four_steps.png', '图5　“侦救一体”四步应急战法', 15.0),
    ('五、双场景比较：四条共性规律', 2): ('fig6_comparison_laws.png', '图6　双场景比较与四条共性规律', 13.5),
}

cur_head, para_idx, in_refs = None, 0, False
i = 0
while i < len(lines):
    ln = lines[i].rstrip()
    i += 1
    if not ln.strip() or ln.strip() == '---':
        continue

    # 公式占位
    m = re.match(r'^\$\$(FORMULA\d)\$\$$', ln.strip())
    if m:
        img, num, w = FORMULAS[m.group(1)]
        add_formula(doc, img, num, w)
        continue

    # 表格
    if ln.startswith('**表') and '|' not in ln:
        cap = re.sub(r'\*\*', '', ln).strip()
        while i < len(lines) and not lines[i].strip():
            i += 1
        rows = []
        while i < len(lines) and lines[i].lstrip().startswith('|'):
            cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
            i += 1
            if all(set(c) <= set('-: ') for c in cells):
                continue
            rows.append(cells)
        if rows:
            add_table(doc, rows, cap)
        continue
    if ln.lstrip().startswith('|'):
        continue

    # 论文标题：二号方正小标宋简体居中
    if ln.startswith('# '):
        p = newp(doc, WD_ALIGN_PARAGRAPH.CENTER, 0, HAO2, before=0, after=8, line=34)
        r = p.add_run(ln[2:].strip()); setfont(r, XBS, HAO2)
        continue
    # 副标题
    if ln.startswith('## ——'):
        p = newp(doc, WD_ALIGN_PARAGRAPH.CENTER, 0, HAO3, before=0, after=12, line=26)
        r = p.add_run(ln[3:].strip()); setfont(r, XBS, HAO3)
        continue
    # 作者姓名：三号楷体_GB2312居中
    if ln.startswith('@@AUTHOR@@'):
        p = newp(doc, WD_ALIGN_PARAGRAPH.CENTER, 0, HAO3, before=0, after=4, line=26)
        r = p.add_run(ln.replace('@@AUTHOR@@', '').strip()); setfont(r, KAI, HAO3)
        continue
    # 单位邮箱：四号宋体居中
    if ln.startswith('@@AFFIL@@'):
        p = newp(doc, WD_ALIGN_PARAGRAPH.CENTER, 0, HAO4, before=0, after=12, line=24)
        r = p.add_run(ln.replace('@@AFFIL@@', '').strip()); setfont(r, SONG, HAO4)
        continue

    # 文末说明
    if ln.startswith('*') and ln.endswith('*') and not ln.startswith('**'):
        p = newp(doc, WD_ALIGN_PARAGRAPH.JUSTIFY, 2, HAO4, before=10, line=24)
        r = p.add_run(ln.strip('*').strip()); setfont(r, KAI, HAO4)
        r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        continue

    # 参考文献条目：五号宋体
    if in_refs and re.match(r'^\[\d+\]', ln):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf = p.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE; pf.line_spacing = 1.2
        pf.space_after = Pt(2); pf.left_indent = Pt(21); pf.first_line_indent = Pt(-21)
        r = p.add_run(ln); setfont(r, SONG, 10.5)
        continue

    # 一级标题（## 一、…）：三号黑体
    if ln.startswith('## '):
        t = ln[3:].strip()
        if t.startswith('参考文献'):
            in_refs = True
        cur_head, para_idx = t, 0
        if t == '摘　要':
            continue  # 摘要标签与内容合并渲染
        p = newp(doc, WD_ALIGN_PARAGRAPH.LEFT, 0, HAO3, before=12, after=6, line=28)
        r = p.add_run(t); setfont(r, HEI, HAO3)
        continue

    # 摘要正文：标签四号黑体 + 内容四号楷体
    if cur_head == '摘　要' and not ln.startswith('**关键词**'):
        p = newp(doc, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, HAO4, before=4, line=24)
        r = p.add_run('摘　要：'); setfont(r, HEI, HAO4)
        rich(p, ln, KAI, HAO4)
        para_idx += 1
        continue
    if ln.startswith('**关键词**'):
        body = ln.split('：', 1)[1].strip() if '：' in ln else ''
        p = newp(doc, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, HAO4, before=4, after=10, line=24)
        r = p.add_run('关键词：'); setfont(r, HEI, HAO4)
        r2 = p.add_run(body); setfont(r2, KAI, HAO4)
        continue

    # 二级标题 **（一）…**：三号楷体
    if re.match(r'^\*\*（[一二三四五六七八九十]+）', ln) and ln.endswith('**'):
        t = ln.strip('*').strip()
        cur_head, para_idx = t, 0
        p = newp(doc, WD_ALIGN_PARAGRAPH.LEFT, 2, HAO3, before=8, after=4, line=28)
        r = p.add_run(t); setfont(r, KAI, HAO3)
        continue

    # 三级标题 **1. …**：三号仿宋_GB2312加粗
    if re.match(r'^\*\*\d+[\.．]', ln) and ln.endswith('**'):
        t = ln.strip('*').strip()
        cur_head, para_idx = t, 0
        p = newp(doc, WD_ALIGN_PARAGRAPH.LEFT, 2, HAO3, before=6, after=3, line=28)
        r = p.add_run(t); setfont(r, FS, HAO3, True)
        continue

    # 正文：三号仿宋_GB2312，首行缩进2字符
    p = newp(doc, WD_ALIGN_PARAGRAPH.JUSTIFY, 2, HAO3, after=0, line=LINE_BODY)
    rich(p, ln, FS, HAO3)
    para_idx += 1

    key = (cur_head, para_idx)
    if key in FIGS_AFTER:
        fn, cap, w = FIGS_AFTER[key]
        add_picture(doc, fn, cap, w)

doc.save(OUT)
print('saved:', OUT)
